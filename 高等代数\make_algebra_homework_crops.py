from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches

BASE = Path(__file__).resolve().parent
PAGE_DIR = BASE / "render_pages"
CROP_DIR = BASE / "question_crops"
CROP_DIR.mkdir(exist_ok=True)

PDF_W, PDF_H = 612.1, 792.0


def page_file(pdfnum: int) -> Path:
    return PAGE_DIR / f"page-{pdfnum:03d}.png"


def crop_segment(pdfnum, y0, y1, x0=85, x1=535):
    im = Image.open(page_file(pdfnum)).convert("RGB")
    sx, sy = im.width / PDF_W, im.height / PDF_H
    box = (
        max(0, int(x0 * sx)),
        max(0, int((y0 - 4) * sy)),
        min(im.width, int(x1 * sx)),
        min(im.height, int(y1 * sy)),
    )
    return im.crop(box)


def make_crop(name, segments):
    pieces = [crop_segment(*seg) for seg in segments]
    if len(pieces) == 1:
        out = pieces[0]
    else:
        width = max(p.width for p in pieces)
        gap = 18
        height = sum(p.height for p in pieces) + gap * (len(pieces) - 1)
        out = Image.new("RGB", (width, height), "white")
        y = 0
        for p in pieces:
            out.paste(p, (0, y))
            y += p.height + gap
    path = CROP_DIR / f"{name}.png"
    out.save(path)
    return path


# Segments are (PDF page number, top, bottom, optional left, optional right).
assignments = {
    "Exercise 1: 10，11": [
        ("C2_E1_Q10", [(56, 169, 240)]),
        ("C2_E1_Q11", [(56, 240, 390)]),
    ],
    "Exercise 2: 35，37": [
        ("C2_E2_Q35", [(73, 69, 248)]),
        ("C2_E2_Q37", [(73, 492, 535)]),
    ],
    "Exercise 3: 6": [
        ("C2_E3_Q6", [(79, 336, 560)]),
    ],
    "Exercise 4: 3": [
        ("C2_E4_Q3", [(86, 77, 138)]),
    ],
    "Exercise 5: 1": [
        ("C2_E5_Q1", [(94, 103, 312)]),
    ],
    "Exercise 6: 1": [
        ("C2_E6_Q1", [(96, 269, 365)]),
    ],
    "Exercise 1: 5": [
        ("C3_E1_Q5", [(102, 484, 610)]),
    ],
    "Exercise 2: 1": [
        ("C3_E2_Q1", [(108, 288, 337)]),
    ],
    "Exercise 3: 4": [
        ("C3_E3_Q4", [(113, 317, 339)]),
    ],
    "Exercise 4: 6": [
        ("C3_E4_Q6", [(119, 71, 121)]),
    ],
    "Exercise 5: 1": [
        ("C3_E5_Q1", [(124, 101, 236)]),
    ],
    "Exercise 6: 3": [
        ("C3_E6_Q3", [(131, 216, 390)]),
    ],
    "Exercise 1: 8": [
        ("C4_E1_Q8", [(136, 231, 266)]),
    ],
    "Exercise 2: 10": [
        ("C4_E2_Q10", [(144, 280, 372)]),
    ],
    "Exercise 3: 11": [
        ("C4_E3_Q11", [(151, 299, 575)]),
    ],
    "Exercise 4: 5": [
        ("C4_E4_Q5", [(158, 519, 640)]),
    ],
    "Exercise 5: 4": [
        ("C4_E5_Q4", [(166, 200, 334)]),
    ],
    "Exercise 1: 4，5": [
        ("C5_E1_Q4", [(172, 72, 150)]),
        ("C5_E1_Q5", [(172, 150, 229)]),
    ],
    "Exercise 2: 6，7": [
        ("C5_E2_Q6", [(178, 462, 522)]),
        ("C5_E2_Q7", [(178, 526, 558)]),
    ],
}

# Duplicate exercise labels occur under different chapters. Disambiguate by current chapter.
chapter_assignments = {
    "Chapter 2 Homework：": {
        "Exercise 1: 10，11": assignments["Exercise 1: 10，11"],
        "Exercise 2: 35，37": assignments["Exercise 2: 35，37"],
        "Exercise 3: 6": assignments["Exercise 3: 6"],
        "Exercise 4: 3": assignments["Exercise 4: 3"],
        "Exercise 5: 1": [("C2_E5_Q1", [(94, 103, 312)])],
        "Exercise 6: 1": assignments["Exercise 6: 1"],
    },
    "Chapter 3 Homework：": {
        "Exercise 1: 5": assignments["Exercise 1: 5"],
        "Exercise 2: 1": assignments["Exercise 2: 1"],
        "Exercise 3: 4": assignments["Exercise 3: 4"],
        "Exercise 4: 6": assignments["Exercise 4: 6"],
        "Exercise 5: 1": [("C3_E5_Q1", [(124, 101, 236)])],
        "Exercise 6: 3": assignments["Exercise 6: 3"],
    },
    "Chapter 4 Homework：": {
        "Exercise 1: 8": assignments["Exercise 1: 8"],
        "Exercise 2: 10": assignments["Exercise 2: 10"],
        "Exercise 3: 11": assignments["Exercise 3: 11"],
        "Exercise 4: 5": assignments["Exercise 4: 5"],
        "Exercise 5: 4": assignments["Exercise 5: 4"],
    },
    "Chapter 5 Homework：": {
        "Exercise 1: 4，5": assignments["Exercise 1: 4，5"],
        "Exercise 2: 6，7": assignments["Exercise 2: 6，7"],
    },
}

generated = {}
for chapter, rows in chapter_assignments.items():
    for row, items in rows.items():
        generated[(chapter, row)] = [(name, make_crop(name, segs)) for name, segs in items]

doc = Document(str(BASE / "高代1作业.docx"))
current_chapter = None
for para in list(doc.paragraphs):
    text = para.text.strip()
    if text in chapter_assignments:
        current_chapter = text
        continue
    key = (current_chapter, text)
    if key not in generated:
        continue
    anchor = para
    for _name, img in generated[key]:
        p = doc.add_paragraph()
        p.add_run().add_picture(str(img), width=Inches(5.75))
        anchor._p.addnext(p._p)
        anchor = p

out = BASE / "高代1作业逐题截图.docx"
doc.save(str(out))
print(out)
print("crops", sum(len(v) for v in generated.values()))
