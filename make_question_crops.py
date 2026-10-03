from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches

ROOT = Path(".")
PAGE_DIR = ROOT / "render_pages"
CROP_DIR = ROOT / "question_crops"
CROP_DIR.mkdir(exist_ok=True)

PAGE_W, PAGE_H = 552.96, 731.76


def col_box(printed, col):
    if col == "full":
        return 8, 545
    if col == "leftwide":
        return 30, 302
    if col == "rightclean":
        return 302, 548
    if printed % 2 == 0:
        return (30, 278) if col == "left" else (280, 548)
    return (8, 252) if col == "left" else (258, 545)


def crop_segment(printed, col, y0, y1):
    image_path = PAGE_DIR / f"page-{printed + 4}.png"
    im = Image.open(image_path).convert("RGB")
    sx = im.width / PAGE_W
    sy = im.height / PAGE_H
    x0, x1 = col_box(printed, col)
    pad_x, top_pad, bottom_pad = 0, 5, 0
    box = (
        max(0, int((x0 - pad_x) * sx)),
        max(0, int((y0 - top_pad) * sy)),
        min(im.width, int((x1 + pad_x) * sx)),
        min(im.height, int((y1 + bottom_pad) * sy)),
    )
    return im.crop(box)


def make_crop(name, segments):
    pieces = [crop_segment(*seg) for seg in segments]
    if len(pieces) == 1:
        out = pieces[0]
    else:
        width = max(p.width for p in pieces)
        gap = 14
        height = sum(p.height for p in pieces) + gap * (len(pieces) - 1)
        out = Image.new("RGB", (width, height), "white")
        y = 0
        for p in pieces:
            out.paste(p, (0, y))
            y += p.height + gap
    path = CROP_DIR / f"{name}.png"
    out.save(path)
    return path


specs = {
    "L5": [
        ("L5_Problems_1", [(227, "left", 63, 162)]),
        ("L5_Problems_7", [(227, "right", 198, 278)]),
        ("L5_Problems_13", [(228, "left", 136, 239)]),
        ("L5_Problems_14", [(228, "left", 239, 280)]),
        ("L5_Problems_15", [(228, "left", 280, 381)]),
        ("L5_Problems_33", [(229, "left", 472, 525)]),
        ("L5_Problems_40", [(229, "right", 393, 422)]),
        ("L5_Theoretical_9", [(230, "right", 203, 339)]),
        ("L5_Theoretical_20", [(231, "left", 172, 245)]),
        ("L5_SelfTest_2", [(232, "left", 278, 359)]),
        ("L5_SelfTest_4", [(232, "left", 421, 531)]),
        ("L5_SelfTest_10", [(232, "right", 455, 590)]),
        ("L5_SelfTest_13", [(233, "left", 339, 413)]),
        ("L5_SelfTest_16", [(233, "right", 40, 136)]),
    ],
    "L6": [
        ("L6_Problems_1", [(290, "leftwide", 67, 169)]),
        ("L6_Problems_8", [(290, "leftwide", 517, 623)]),
        ("L6_Problems_9", [(290, "leftwide", 623, 690), (290, "rightclean", 60, 146)]),
        ("L6_Problems_10", [(290, "rightclean", 153, 225)]),
        ("L6_Problems_15", [(290, "rightclean", 560, 690), (291, "left", 40, 104)]),
        ("L6_Problems_19", [(291, "left", 521, 625)]),
        ("L6_Problems_20", [(291, "left", 625, 690), (291, "right", 40, 122)]),
        ("L6_Problems_40", [(292, "rightclean", 531, 680)]),
        ("L6_Problems_41", [(293, "left", 39, 130)]),
        ("L6_Theoretical_6", [(294, "left", 423, 506)]),
        ("L6_Theoretical_25", [(296, "left", 39, 134)]),
        ("L6_SelfTest_2", [(296, "left", 642, 690), (296, "right", 436, 449)]),
        ("L6_SelfTest_6", [(297, "left", 95, 239)]),
        ("L6_SelfTest_7", [(297, "left", 239, 506)]),
    ],
    "L7": [
        ("L7_Problems_1", [(375, "left", 234, 286)]),
        ("L7_Problems_4", [(375, "left", 544, 652)]),
        ("L7_Problems_6", [(375, "right", 308, 338)]),
        ("L7_Problems_30", [(377, "right", 375, 438)]),
        ("L7_Problems_33", [(377, "right", 500, 546)]),
        ("L7_Problems_40", [(378, "left", 242, 316)]),
        ("L7_Theoretical_1", [(382, "left", 69, 88)]),
        ("L7_Theoretical_19", [(383, "right", 140, 198)]),
        ("L7_Theoretical_22", [(383, "right", 506, 566)]),
        ("L7_SelfTest_32", [(389, "right", 40, 178)]),
    ],
    "L8": [
        ("L8_Problems_1", [(420, "left", 430, 468)]),
        ("L8_Problems_4", [(420, "left", 650, 690), (420, "rightclean", 430, 558)]),
        ("L8_Problems_13", [(421, "left", 587, 690), (421, "right", 40, 115)]),
        ("L8_Problems_15", [(421, "right", 200, 252)]),
        ("L8_Theoretical_1", [(422, "left", 248, 314)]),
        ("L8_Theoretical_2", [(422, "left", 314, 423)]),
        ("L8_SelfTest_1", [(423, "left", 350, 419)]),
        ("L8_SelfTest_2", [(423, "left", 419, 499)]),
        ("L8_SelfTest_8", [(424, "left", 68, 120)]),
    ],
}

generated = {section: [(name, make_crop(name, segs)) for name, segs in items] for section, items in specs.items()}

doc = Document("新建 DOCX 文档.docx")
headings = {p.text.strip(): p for p in doc.paragraphs if p.text.strip() in generated}
for heading in ["L8", "L7", "L6", "L5"]:
    anchor = headings[heading]
    for _name, img_path in generated[heading]:
        p = doc.add_paragraph()
        p.add_run().add_picture(str(img_path), width=Inches(5.77))
        anchor._p.addnext(p._p)
        anchor = p

out = "概率论作业逐题截图 L1-L8 修正版.docx"
doc.save(out)
print(Path(out).resolve())
print("crops", sum(len(v) for v in generated.values()))
